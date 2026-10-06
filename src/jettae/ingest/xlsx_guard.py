"""Resource limits for XLSX applied *before* openpyxl touches the workbook.

Why a separate pre-scan: openpyxl trusts the file. In normal mode it materialises every cell
of every merged range (one ``<mergeCell ref="A1:XFD1048576"/>`` is 17 billion objects); in
read-only mode it sizes iteration from the ``<dimension>`` tag, which the file can set to
``A1:XFD1048576`` while holding three cells. Even before any cell, ``load_workbook`` builds
Python objects for every element of the content-type manifest, the workbook part, the
relationship parts, ``xl/styles.xml``, the document properties and chart sheets; a
sub-megabyte file with a few hundred thousand ``<xf>`` elements costs seconds and gigabytes.
This module bounds the work from what the file actually contains:

1. Zip container: entry count, duplicate names, per-entry and total uncompressed size,
   compression ratio (per large entry and overall). ``zipfile`` stops decompressing at the
   declared size and checks the CRC, so declared sizes bound real work.
2. Every member is sniffed the way libxml2 and expat detect encodings. XML is accepted only
   as UTF-8 (with or without BOM) or UTF-16 *with* a BOM; UTF-16 without a BOM, UCS-4,
   EBCDIC or any other declared encoding is refused, because the scan below and openpyxl
   must decode the same characters. A DTD (``<!DOCTYPE``, the only place entities can be
   declared) is refused in every XML member.
3. The parts openpyxl will parse are found the way openpyxl finds them - content types for
   the workbook and shared-string parts, the workbook relationships for sheets, fixed paths
   for styles and document properties, and every part reachable from a chart sheet - not by
   the name of their root element, which the file controls. A part that openpyxl would
   parse but that is not scannable XML fails closed.
4. Worksheets are stream-scanned with ``lxml.iterparse`` (no entity resolution, no network,
   no huge-tree mode), counting rows, the largest row/column index, cells, merged ranges and
   merged area. The ``<dimension>`` tag is never read. Shared strings are counted per table
   and per string; every other parsed part has an element cap, styles also a cap on style
   records. Each cap stops the scan as soon as it is crossed.

The caller then iterates exactly the scanned bounds (``SheetScan.max_row/max_col``) and maps
openpyxl's sheets to scanned parts by path; a sheet openpyxl loads that was not scanned is
treated as a corrupt file (fail closed).
"""

from __future__ import annotations

import io
import posixpath
import re
import zipfile
from collections.abc import Callable
from dataclasses import dataclass, field
from typing import IO, Any

from lxml import etree
from openpyxl.xml.constants import (
    ARC_CONTENT_TYPES,
    ARC_CORE,
    ARC_CUSTOM,
    ARC_STYLE,
    ARC_WORKBOOK,
    SHARED_STRINGS,
    XLSM,
    XLSX,
    XLTM,
    XLTX,
)

from jettae.ingest.cells import CorruptFile, IngestRejected
from jettae.ingest.security import Limits

EXCEL_MAX_ROW = 1_048_576
EXCEL_MAX_COL = 16_384
_REF = re.compile(r"^\$?([A-Za-z]{1,3})\$?(\d{1,7})$")
_ROOT = re.compile(r"<([A-Za-z_][\w.\-]*(?::[A-Za-z_][\w.\-]*)?)")
_XML_DECL_ENCODING = re.compile(r"""^<\?xml\s[^>]*?\bencoding\s*=\s*["']([^"']*)["']""")
# same order as openpyxl.reader.excel._find_workbook_part
_WORKBOOK_TYPES = (XLTM, XLTX, XLSM, XLSX)
_MACRO_CT = re.compile(r"macroEnabled|vbaProject|ms-excel\.sheet\.binary|ms-excel\.addin", re.I)
_STYLE_RECORDS = frozenset({"numFmt", "font", "fill", "border", "xf", "dxf", "cellStyle"})
# byte patterns libxml2 (xmlDetectCharEncoding) and expat use to pick a non-UTF-8 decoding
_REFUSED_BOMS = (b"\x00\x00\xfe\xff", b"\xff\xfe\x00\x00")  # UTF-32; before the UTF-16 BOMs
_REFUSED_HEADS = (
    b"\x00\x00\x00<",
    b"<\x00\x00\x00",
    b"\x00\x00<\x00",
    b"\x00<\x00\x00",
    b"\x4c\x6f\xa7\x94",  # EBCDIC "<?xm"
)


@dataclass(frozen=True)
class XlsxLimits:
    """Caps for one workbook. Defaults follow :class:`~jettae.ingest.security.Limits`
    (rows, columns, sheets, zip sizes) plus XLSX-specific caps. ``max_cell_area`` bounds the
    dense grid the parser builds (sum over sheets of rows x columns), which is what makes
    a sparse file with one far-away cell expensive. The part caps bound what openpyxl
    materialises outside the cell grid (about 1-2 KiB of Python objects per element)."""

    max_rows: int = 100_000  # sum of the largest row index of each sheet
    max_cols: int = 200
    max_sheets: int = 50
    max_cells: int = 1_000_000  # <c> elements in all sheets
    max_cell_area: int = 1_000_000  # sum of max_row * max_col over sheets
    max_merged_ranges: int = 10_000
    max_merged_area: int = 1_000_000  # cells covered by merged ranges, all sheets
    max_row_elements: int = 200_000  # <row> elements in all sheets
    max_shared_strings: int = 1_000_000  # <si> elements
    max_shared_string_elements: int = 4_000_000  # all elements of the shared-string part
    max_string_elements: int = 10_000  # elements inside one <si> (rich-text runs)
    # Excel allows 64 000 cell formats; bloated real workbooks get close to that
    max_style_records: int = 100_000  # numFmt + font + fill + border + xf + dxf + cellStyle
    max_part_elements: int = 250_000  # any other part openpyxl parses (manifest, rels, ...)
    max_aux_elements: int = 750_000  # sum over those parts
    max_zip_entries: int = 5000
    max_uncompressed: int = 200 * 1024 * 1024
    max_entry_uncompressed: int = 128 * 1024 * 1024
    max_ratio: int = 200
    max_prolog_bytes: int = 64 * 1024

    @classmethod
    def from_limits(cls, limits: Limits, **overrides: Any) -> XlsxLimits:
        base: dict[str, Any] = {
            "max_rows": limits.max_rows,
            "max_cols": limits.max_cols,
            "max_sheets": limits.max_sheets,
            "max_zip_entries": limits.max_zip_entries,
            "max_uncompressed": limits.max_uncompressed,
            "max_ratio": limits.max_ratio,
            "max_row_elements": max(2 * limits.max_rows, 1),
        }
        base.update(overrides)
        return cls(**base)


@dataclass
class SheetScan:
    part: str
    max_row: int = 0
    max_col: int = 0
    cells: int = 0
    rows: int = 0
    merged: list[tuple[int, int, int, int]] = field(default_factory=list)  # r0, c0, r1, c1
    has_formula: bool = False

    @property
    def area(self) -> int:
        return self.max_row * self.max_col


@dataclass
class WorkbookScan:
    sheets: dict[str, SheetScan]  # zip part name -> scan
    has_formula: bool
    shared_strings: int
    sheet_entries: int  # <sheet> elements in the workbook part(s)
    aux_elements: int = 0  # elements in the other parts openpyxl parses


def col_index(letters: str) -> int:
    n = 0
    for ch in letters.upper():
        n = n * 26 + (ord(ch) - 64)
    return n


def parse_ref(ref: str) -> tuple[int, int]:
    """``"B12"`` -> (12, 2). Out-of-grid or malformed references make the file corrupt."""
    m = _REF.match(ref.strip())
    if not m:
        raise CorruptFile(f"invalid cell reference {ref[:32]!r}", code="bad_xlsx")
    row, col = int(m.group(2)), col_index(m.group(1))
    if not (1 <= row <= EXCEL_MAX_ROW and 1 <= col <= EXCEL_MAX_COL):
        raise CorruptFile(f"cell reference {ref[:32]!r} is outside the sheet", code="bad_xlsx")
    return row, col


def _parse_row_number(raw: str) -> int:
    # openpyxl accepts integral floats ("5.0") for row numbers; mirror it
    try:
        n = int(raw)
    except ValueError:
        try:
            f = float(raw)
        except ValueError:
            raise CorruptFile(f"invalid row number {raw[:16]!r}", code="bad_xlsx") from None
        if not f.is_integer():
            raise CorruptFile(f"invalid row number {raw[:16]!r}", code="bad_xlsx") from None
        n = int(f)
    if not 1 <= n <= EXCEL_MAX_ROW:
        raise CorruptFile(f"row number {n} is outside the sheet", code="bad_xlsx")
    return n


# ------------------------------------------------------------------------------- zip level
def check_container(zf: zipfile.ZipFile, lim: XlsxLimits) -> None:
    infos = zf.infolist()
    if len(infos) > lim.max_zip_entries:
        raise IngestRejected(f"{len(infos)} zip entries exceed limit", code="zip_entries")
    names: set[str] = set()
    total = total_compressed = 0
    for info in infos:
        if info.filename in names:
            # two members with one name: readers may disagree on which one is "the" part
            raise CorruptFile(f"duplicate zip member {info.filename!r}", code="bad_zip")
        names.add(info.filename)
        if info.file_size > lim.max_entry_uncompressed:
            raise IngestRejected(
                f"zip member {info.filename!r} expands to {info.file_size} bytes",
                code="zip_entry_size",
            )
        if info.file_size > 1024 * 1024 and (
            info.file_size / max(info.compress_size, 1) > lim.max_ratio
        ):
            raise IngestRejected(
                f"suspicious compression ratio in {info.filename}", code="zip_bomb"
            )
        total += info.file_size
        total_compressed += info.compress_size
    if total > lim.max_uncompressed:
        raise IngestRejected("uncompressed size exceeds limit", code="zip_bomb")
    if total > 1024 * 1024 and total / max(total_compressed, 1) > lim.max_ratio:
        raise IngestRejected("suspicious overall compression ratio", code="zip_bomb")


# ------------------------------------------------------------------------- member sniffing
def _refuse_encoding(name: str, what: str) -> IngestRejected:
    return IngestRejected(
        f"XML part {name!r} uses {what}; only UTF-8 or UTF-16 with a byte order mark is accepted",
        code="xml_encoding",
    )


def _decode_head(name: str, head: bytes) -> tuple[str, str] | None:
    """Text of the first bytes and the encoding family a parser would use, or None when the
    member cannot be XML for libxml2 or expat. Raises for XML in a refused encoding."""
    if head.startswith(_REFUSED_BOMS):
        raise _refuse_encoding(name, "a UTF-32 byte order mark")
    if head.startswith(b"\xef\xbb\xbf"):
        return head[3:].decode("utf-8", errors="ignore"), "utf-8"
    if head.startswith((b"\xff\xfe", b"\xfe\xff")):
        return head.decode("utf-16", errors="ignore").lstrip("﻿"), "utf-16"
    if head.startswith(_REFUSED_HEADS):
        raise _refuse_encoding(name, "UCS-4 or EBCDIC without a byte order mark")
    if len(head) >= 2 and (head[0] == 0 or head[1] == 0):
        # expat reads a NUL in either of the first two bytes as UTF-16 without a BOM (and
        # libxml2 does for "<?"); a binary member (EMF image, OLE blob) is not XML in it
        codec = "utf-16-be" if head[0] == 0 else "utf-16-le"
        text = head[: len(head) - len(head) % 2].decode(codec, errors="ignore")
        if text.lstrip(" \t\r\n").startswith("<"):
            raise _refuse_encoding(name, "UTF-16 without a byte order mark")
        return None
    return head.decode("utf-8", errors="ignore"), "utf-8"


def xml_root(zf: zipfile.ZipFile, name: str, lim: XlsxLimits) -> str | None:
    """Local name of the root element of ``name``, or None when the member is not XML.
    Raises when the member is XML in a refused encoding, when its prolog holds a DTD or
    is longer than ``max_prolog_bytes``."""
    with zf.open(name) as fh:
        head = fh.read(lim.max_prolog_bytes)
    truncated = len(head) == lim.max_prolog_bytes
    decoded = _decode_head(name, head)
    if decoded is None:
        return None
    text, family = decoded
    pos, n = 0, len(text)
    first = True
    while True:
        while pos < n and text[pos] in " \t\r\n":
            pos += 1
        if pos >= n:
            if truncated:
                raise IngestRejected(f"XML prolog of {name!r} too long", code="xml_prolog")
            return None
        if text.startswith("<?", pos):
            end = text.find("?>", pos)
            if first and text.startswith("<?xml", pos) and end >= 0:
                m = _XML_DECL_ENCODING.match(text, pos, end + 2)
                declared = m.group(1).strip().lower() if m else ""
                allowed = {"", "utf-8", "utf8"} if family == "utf-8" else {"", "utf-16"}
                if declared not in allowed and not (
                    family == "utf-16" and declared in ("utf-16le", "utf-16be")
                ):
                    raise _refuse_encoding(name, f"the declared encoding {declared[:20]!r}")
        elif text.startswith("<!--", pos):
            end = text.find("-->", pos)
        elif text.startswith("<!", pos):
            raise IngestRejected(
                f"XML part {name!r} declares a DTD; entity declarations are not accepted",
                code="xml_dtd",
            )
        elif text.startswith("<", pos):
            m = _ROOT.match(text, pos)
            if not m:
                return None
            return m.group(1).rsplit(":", 1)[-1]
        else:
            return None  # binary member (image, printer settings, ...)
        first = False
        if end < 0:
            if truncated:
                raise IngestRejected(f"XML prolog of {name!r} too long", code="xml_prolog")
            return None
        pos = end + 2


def _iterparse(fh: IO[bytes], events: tuple[str, ...], tags: tuple[str, ...] | None) -> Any:
    return etree.iterparse(
        fh,
        events=events,
        tag=tags,
        resolve_entities=False,
        no_network=True,
        load_dtd=False,
        huge_tree=False,
        recover=False,
        remove_comments=True,
    )


def _drop_done(el: Any) -> None:
    """Free a processed element and its already processed siblings (bounded memory)."""
    el.clear(keep_tail=False)
    parent = el.getparent()
    if parent is not None:
        while el.getprevious() is not None:
            del parent[0]


def _local(el: Any) -> str:
    tag = el.tag
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) else ""


class _Budget:
    """Workbook-wide counters shared by the part scans; every increment checks its cap."""

    def __init__(self, lim: XlsxLimits) -> None:
        self.lim = lim
        self.rows_before = 0  # sum of max_row of sheets already scanned
        self.area_before = 0
        self.cells = 0
        self.row_elements = 0
        self.merged_ranges = 0
        self.merged_area = 0
        self.aux_elements = 0

    def extent(self, sc: SheetScan, row: int, col: int) -> None:
        lim = self.lim
        if row > sc.max_row:
            sc.max_row = row
            if self.rows_before + row > lim.max_rows:
                raise IngestRejected(
                    f"rows exceed limit {lim.max_rows} (row {row} in {sc.part})", code="rows"
                )
        if col > sc.max_col:
            sc.max_col = col
            if col > lim.max_cols:
                raise IngestRejected(
                    f"{sc.part}: column {col} exceeds limit {lim.max_cols}", code="cols"
                )
        if self.area_before + sc.max_row * sc.max_col > lim.max_cell_area:
            raise IngestRejected(
                f"sheet grid {sc.max_row}x{sc.max_col} exceeds {lim.max_cell_area} cells",
                code="cell_area",
            )


def scan_sheet(fh: IO[bytes], part: str, budget: _Budget) -> SheetScan:
    lim = budget.lim
    sc = SheetScan(part)
    row_counter = col_counter = 0
    try:
        for event, el in _iterparse(
            fh, ("start", "end"), ("{*}row", "{*}c", "{*}mergeCell", "{*}f")
        ):
            tag = _local(el)
            if tag == "c":
                if event == "start":
                    ref = el.get("r")
                    if ref:
                        row, col = parse_ref(ref)
                    else:
                        row, col = row_counter, col_counter + 1
                        if row < 1 or col > EXCEL_MAX_COL:
                            raise CorruptFile(f"cell outside the sheet in {part}", code="bad_xlsx")
                    col_counter = col
                    sc.cells += 1
                    budget.cells += 1
                    if budget.cells > lim.max_cells:
                        raise IngestRejected(f"cells exceed limit {lim.max_cells}", code="cells")
                    budget.extent(sc, row, col)
                else:
                    _drop_done(el)
            elif tag == "row":
                if event == "start":
                    raw = el.get("r")
                    row_counter = _parse_row_number(raw) if raw else row_counter + 1
                    if row_counter > EXCEL_MAX_ROW:
                        raise CorruptFile(f"row outside the sheet in {part}", code="bad_xlsx")
                    col_counter = 0
                    sc.rows += 1
                    budget.row_elements += 1
                    if budget.row_elements > lim.max_row_elements:
                        raise IngestRejected(
                            f"row elements exceed limit {lim.max_row_elements}", code="rows"
                        )
                else:
                    _drop_done(el)
            elif tag == "mergeCell" and event == "end":
                ref = el.get("ref") or ""
                a, _, b = ref.partition(":")
                r0, c0 = parse_ref(a)
                r1, c1 = parse_ref(b) if b else (r0, c0)
                r0, r1 = sorted((r0, r1))
                c0, c1 = sorted((c0, c1))
                budget.merged_ranges += 1
                if budget.merged_ranges > lim.max_merged_ranges:
                    raise IngestRejected(
                        f"merged ranges exceed limit {lim.max_merged_ranges}",
                        code="merged_ranges",
                    )
                budget.merged_area += (r1 - r0 + 1) * (c1 - c0 + 1)
                if budget.merged_area > lim.max_merged_area:
                    raise IngestRejected(
                        f"merged area exceeds limit {lim.max_merged_area} cells",
                        code="merged_area",
                    )
                # openpyxl (normal mode) counts merged ranges in a sheet's size; so do we
                budget.extent(sc, r1, c1)
                sc.merged.append((r0, c0, r1, c1))
                _drop_done(el)
            elif tag == "f" and event == "end":
                sc.has_formula = True
    except etree.XMLSyntaxError as e:
        raise CorruptFile(f"malformed worksheet XML in {part}: {e}", code="bad_xlsx") from e
    budget.rows_before += sc.max_row
    budget.area_before += sc.area
    return sc


def _walk(
    fh: IO[bytes],
    part: str,
    cap: int,
    on_start: Callable[[str], None] | None = None,
    on_end: Callable[[str, Any], None] | None = None,
) -> int:
    """Stream every element of ``part``; raise as soon as there are more than ``cap``."""
    n = 0
    try:
        for event, el in _iterparse(fh, ("start", "end"), None):
            name = _local(el)
            if event == "start":
                n += 1
                if n > cap:
                    raise IngestRejected(
                        f"{part} has more than {cap} XML elements", code="xml_elements"
                    )
                if on_start is not None:
                    on_start(name)
            else:
                if on_end is not None:
                    on_end(name, el)
                _drop_done(el)
    except etree.XMLSyntaxError as e:
        raise CorruptFile(f"malformed XML in {part}: {e}", code="bad_xlsx") from e
    return n


def _scan_aux(
    zf: zipfile.ZipFile,
    part: str,
    budget: _Budget,
    on_start: Callable[[str], None] | None = None,
    on_end: Callable[[str, Any], None] | None = None,
) -> None:
    lim = budget.lim
    cap = min(lim.max_part_elements, lim.max_aux_elements - budget.aux_elements)
    with zf.open(part) as fh:
        try:
            budget.aux_elements += _walk(fh, part, max(cap, 0), on_start, on_end)
        except IngestRejected as e:
            if e.code == "xml_elements" and cap < lim.max_part_elements:
                raise IngestRejected(
                    f"XML parts openpyxl parses exceed {lim.max_aux_elements} elements",
                    code="xml_elements",
                ) from e
            raise


def _scan_shared_strings(zf: zipfile.ZipFile, part: str, lim: XlsxLimits) -> int:
    count = 0
    inside = 0  # elements of the current <si>, which openpyxl builds into one object

    def start(name: str) -> None:
        nonlocal count, inside
        if name == "si":
            count += 1
            inside = 0
            if count > lim.max_shared_strings:
                raise IngestRejected(
                    f"si elements in {part} exceed limit {lim.max_shared_strings}",
                    code="shared_strings",
                )
            return
        inside += 1
        if inside > lim.max_string_elements:
            raise IngestRejected(
                f"one shared string in {part} has more than {lim.max_string_elements} elements",
                code="shared_strings",
            )

    with zf.open(part) as fh:
        try:
            _walk(fh, part, lim.max_shared_string_elements, start)
        except IngestRejected as e:
            if e.code == "xml_elements":
                raise IngestRejected(str(e), code="shared_strings") from e
            raise
    return count


# ---------------------------------------------------------------------- package resolution
def _rels_path(part: str) -> str:
    """openpyxl.packaging.relationship.get_rels_path"""
    folder, obj = posixpath.split(part)
    return posixpath.join(folder, "_rels", f"{obj}.rels")


def _rel_target(rels_part: str, target: str) -> str:
    """Target normalisation of openpyxl.packaging.relationship.get_dependents."""
    if target.startswith("/"):
        return target[1:]
    parent = posixpath.split(posixpath.dirname(rels_part))[0]
    return posixpath.normpath(posixpath.join(parent, target))


@dataclass
class _Package:
    roots: dict[str, str]  # XML members -> local root name
    names: set[str]
    budget: _Budget
    zf: zipfile.ZipFile
    scanned: set[str] = field(default_factory=set)

    def require_xml(self, part: str) -> bool:
        """True when ``part`` exists and is scannable XML; False when it does not exist.
        openpyxl would parse an existing part, so one we cannot scan fails closed."""
        if part not in self.names:
            return False
        if part not in self.roots:
            raise CorruptFile(
                f"{part!r} is referenced by the package but is not XML", code="bad_xlsx"
            )
        return True

    def aux(
        self,
        part: str,
        on_start: Callable[[str], None] | None = None,
        on_end: Callable[[str, Any], None] | None = None,
    ) -> bool:
        if part in self.scanned or not self.require_xml(part):
            return part in self.scanned
        self.scanned.add(part)
        _scan_aux(self.zf, part, self.budget, on_start, on_end)
        return True

    def relationships(self, owner: str) -> list[tuple[str, str]]:
        """(type, normalised target) of the internal relationships of ``owner``."""
        rels_part = _rels_path(owner)
        found: list[tuple[str, str]] = []

        def end(name: str, el: Any) -> None:
            if name == "Relationship" and el.get("TargetMode") != "External":
                target = el.get("Target")
                if target:
                    found.append((el.get("Type") or "", _rel_target(rels_part, target)))

        if rels_part in self.scanned:
            # already counted; read the targets again without charging the budget twice
            with self.zf.open(rels_part) as fh:
                _walk(fh, rels_part, self.budget.lim.max_part_elements, None, end)
            return found
        self.aux(rels_part, on_end=end)
        return found


@dataclass
class _Resolved:
    shared_strings: list[str]
    worksheets: list[str]
    sheet_entries: int
    budget: _Budget


def _resolve_and_scan(zf: zipfile.ZipFile, roots: dict[str, str], lim: XlsxLimits) -> _Resolved:
    """Scan the non-worksheet parts openpyxl parses and return the shared-string and
    worksheet parts, chosen the way openpyxl chooses them (plus parts whose root element
    claims the role, so a renamed or extra part is never left unbounded)."""
    budget = _Budget(lim)
    pkg = _Package(roots, set(zf.namelist()), budget, zf)
    overrides: list[tuple[str, str]] = []
    defaults: list[str] = []

    def manifest_end(name: str, el: Any) -> None:
        ct = el.get("ContentType") or ""
        if _MACRO_CT.search(ct):
            raise IngestRejected("macro-enabled or binary workbook is not accepted", code="macro")
        if name == "Override":
            overrides.append(((el.get("PartName") or "").lstrip("/"), ct))
        elif name == "Default":
            defaults.append(ct)

    if not pkg.aux(ARC_CONTENT_TYPES, on_end=manifest_end):
        raise CorruptFile("missing [Content_Types].xml", code="bad_ooxml")

    # openpyxl takes the first matching Override; scanning every candidate (and every part
    # whose root claims the role) is a superset of what it can pick
    workbooks = [p for t in _WORKBOOK_TYPES for p, ct in overrides if ct == t]
    if set(defaults) & set(_WORKBOOK_TYPES):
        workbooks.append(ARC_WORKBOOK)
    workbooks += [n for n, r in roots.items() if r == "workbook"]
    shared = [p for p, ct in overrides if ct == SHARED_STRINGS]
    shared += [n for n, r in roots.items() if r == "sst"]

    sheet_entries = 0

    def workbook_start(name: str) -> None:
        nonlocal sheet_entries
        if name == "sheet":
            sheet_entries += 1
            if sheet_entries > lim.max_sheets:
                raise IngestRejected(
                    f"{sheet_entries} sheet entries exceed limit {lim.max_sheets}", code="sheets"
                )

    worksheets: list[str] = [n for n, r in roots.items() if r in ("worksheet", "dialogsheet")]
    chart_roots: list[str] = []
    workbook_parts: list[str] = []
    for wb_part in dict.fromkeys(workbooks):
        if not pkg.aux(wb_part, on_start=workbook_start):
            continue
        workbook_parts.append(wb_part)
        for rel_type, target in pkg.relationships(wb_part):
            if target not in pkg.names:
                continue  # openpyxl skips sheets whose part is missing
            if "chartsheet" in rel_type:
                chart_roots.append(target)
            elif rel_type.rsplit("/", 1)[-1] in ("worksheet", "dialogsheet", "macrosheet"):
                worksheets.append(target)
            # openpyxl parses the relationships of every sheet it loads, even read-only
            pkg.relationships(target)
    if not workbook_parts:
        raise CorruptFile("file contains no valid workbook part", code="bad_xlsx")

    for fixed in (ARC_CORE, ARC_CUSTOM):
        pkg.aux(fixed)
    style_records = 0

    def style_start(name: str) -> None:
        nonlocal style_records
        if name in _STYLE_RECORDS:
            style_records += 1
            if style_records > lim.max_style_records:
                raise IngestRejected(
                    f"{ARC_STYLE} has more than {lim.max_style_records} style records",
                    code="styles",
                )

    pkg.aux(ARC_STYLE, on_start=style_start)

    # chart sheets are parsed whole, with their drawings and charts: bound everything
    # reachable from them (images and other non-XML targets are only opened, not parsed)
    todo = list(chart_roots)
    while todo:
        part = todo.pop()
        if part in pkg.scanned or part not in pkg.roots:
            continue
        pkg.aux(part)
        todo.extend(t for _, t in pkg.relationships(part) if t in pkg.names)

    shared = list(dict.fromkeys(shared))
    for part in shared:
        pkg.require_xml(part)
    return _Resolved(shared, list(dict.fromkeys(worksheets)), sheet_entries, budget)


def inspect_xlsx(content: bytes, lim: XlsxLimits) -> WorkbookScan:
    """Check the container, every parsed non-cell part and stream-scan every worksheet.
    Raises :class:`IngestRejected` (limit) or :class:`CorruptFile` (malformed)."""
    try:
        zf = zipfile.ZipFile(io.BytesIO(content))
    except zipfile.BadZipFile as e:
        raise CorruptFile(f"not a valid zip/xlsx container: {e}", code="bad_zip") from e
    try:
        with zf:
            check_container(zf, lim)
            roots: dict[str, str] = {}
            for info in zf.infolist():
                if info.is_dir():
                    continue
                root = xml_root(zf, info.filename, lim)
                if root is not None:
                    roots[info.filename] = root
            res = _resolve_and_scan(zf, roots, lim)
            if len(res.worksheets) > lim.max_sheets:
                raise IngestRejected(f"{len(res.worksheets)} sheets exceed limit", code="sheets")
            shared = sum(_scan_shared_strings(zf, p, lim) for p in res.shared_strings)
            scans: dict[str, SheetScan] = {}
            for name in res.worksheets:
                if name not in roots:
                    raise CorruptFile(f"worksheet {name!r} is not XML", code="bad_xlsx")
                with zf.open(name) as fh:
                    scans[name] = scan_sheet(fh, name, res.budget)
    except (zipfile.BadZipFile, zipfile.LargeZipFile, EOFError, NotImplementedError) as e:
        # CRC mismatch, truncated member, encrypted or unsupported compression
        raise CorruptFile(f"damaged zip member: {e}", code="bad_zip") from e
    return WorkbookScan(
        sheets=scans,
        has_formula=any(s.has_formula for s in scans.values()),
        shared_strings=shared,
        sheet_entries=res.sheet_entries,
        aux_elements=res.budget.aux_elements,
    )
