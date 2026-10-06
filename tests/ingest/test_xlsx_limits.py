"""XLSX resource limits are enforced before openpyxl iterates any cell.

Each fixture is a tiny hand-written workbook (not evaluation data) whose XML claims or
contains something expensive. Every case has a wall-clock bound: without the guard these
files take minutes to hours or exhaust memory."""

from __future__ import annotations

import io
import time
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from jettae.domain.status import DocumentStatus
from jettae.ingest import IngestRejected, parse_bytes, safe_parse
from jettae.ingest.cells import CorruptFile
from jettae.ingest.xlsx import parse_xlsx
from jettae.ingest.xlsx_guard import XlsxLimits

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
SHEET = "xl/worksheets/sheet1.xml"
BOUND_S = 5.0  # generous for a slow CI runner; the guarded paths take well under 1 s here


def _base() -> bytes:
    wb = Workbook()
    wb.active["A1"] = "x"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _replace(xlsx: bytes, members: dict[str, bytes]) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(xlsx))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            if item.filename not in members:
                dst.writestr(item, src.read(item.filename))
        for name, data in members.items():
            dst.writestr(name, data)
    return out.getvalue()


def _sheet(sheet_data: str, *, dimension: str = "A1", merges: str = "") -> bytes:
    return (
        f'<worksheet xmlns="{NS}"><dimension ref="{dimension}"/>'
        f"<sheetData>{sheet_data}</sheetData>{merges}</worksheet>"
    ).encode()


def _workbook(sheet_xml: bytes, **extra: bytes) -> bytes:
    return _replace(_base(), {SHEET: sheet_xml, **extra})


@pytest.fixture
def no_openpyxl(monkeypatch):
    """Fail the test if the parser reaches openpyxl."""

    def boom(*a, **kw):
        raise AssertionError("openpyxl load_workbook reached before the limit check")

    monkeypatch.setattr("jettae.ingest.xlsx.load_workbook", boom)


def _rejected(content: bytes, code: str, **kw) -> None:
    t0 = time.perf_counter()
    with pytest.raises(IngestRejected) as e:
        parse_xlsx(content, filename="t.xlsx", **kw)
    assert e.value.code == code, (e.value.code, str(e.value))
    assert time.perf_counter() - t0 < BOUND_S


# ------------------------------------------------------------------ dimension is not trusted
def test_tiny_file_with_huge_dimension_claim_parses_quickly():
    cells = (
        '<row r="1"><c r="A1" t="inlineStr"><is><t>날짜</t></is></c>'
        '<c r="B1" t="inlineStr"><is><t>금액</t></is></c></row>'
        '<row r="2"><c r="A2" t="n"><v>1000</v></c></row>'
    )
    content = _workbook(_sheet(cells, dimension="A1:XFD1048576"))
    assert len(content) < 10_000
    # read-only openpyxl would size its iteration from this claim
    ro = load_workbook(io.BytesIO(content), read_only=True)
    assert (ro.active.max_row, ro.active.max_column) == (1_048_576, 16_384)
    ro.close()
    t0 = time.perf_counter()
    doc = parse_xlsx(content, filename="dim.xlsx")
    assert time.perf_counter() - t0 < BOUND_S
    (table,) = doc.tables
    assert len(table.rows) == 2 and all(len(r.cells) == 2 for r in table.rows)
    assert table.rows[0].texts() == ["날짜", "금액"]
    assert table.rows[1].cells[0].raw == 1000


# ------------------------------------------------------------------ row / column / area caps
def test_single_cell_at_last_excel_cell_is_rejected_before_openpyxl(no_openpyxl):
    content = _workbook(_sheet('<row r="1048576"><c r="XFD1048576"><v>1</v></c></row>'))
    assert len(content) < 10_000
    _rejected(content, "rows")
    _rejected(content, "cols", xlsx_limits=XlsxLimits(max_rows=2_000_000, max_cell_area=10**12))


def test_sparse_far_cell_exceeds_grid_area(no_openpyxl):
    # within the row (100 000) and column (200) caps, but the dense grid is ~20 million cells
    cells = (
        '<row r="1"><c r="A1"><v>1</v></c></row><row r="99999"><c r="GR99999"><v>2</v></c></row>'
    )
    _rejected(_workbook(_sheet(cells)), "cell_area")


def test_too_many_cells(no_openpyxl):
    rows = "".join(
        f'<row r="{r}">' + "".join(f'<c r="{c}{r}"><v>1</v></c>' for c in "ABCDE") + "</row>"
        for r in range(1, 301)
    )
    _rejected(_workbook(_sheet(rows)), "cells", xlsx_limits=XlsxLimits(max_cells=1000))


def test_too_many_row_elements(no_openpyxl):
    # no cells at all (so no row/column extent), yet 200 001 row elements to walk
    rows = "".join(f'<row r="{r}"/>' for r in range(1, 200_002))
    _rejected(_workbook(_sheet(rows)), "rows")


# ------------------------------------------------------------------ merged cells
def test_one_merged_range_over_the_whole_sheet(no_openpyxl):
    sheet = _sheet(
        '<row r="1"><c r="A1"><v>1</v></c></row>',
        merges='<mergeCells count="1"><mergeCell ref="A1:XFD1048576"/></mergeCells>',
    )
    content = _workbook(sheet)
    assert len(content) < 10_000
    _rejected(content, "merged_area")


def test_too_many_merged_ranges(no_openpyxl):
    merges = "".join(f'<mergeCell ref="A{r}:B{r}"/>' for r in range(1, 10_002))
    sheet = _sheet(
        '<row r="1"><c r="A1"><v>1</v></c></row>',
        merges=f'<mergeCells count="10001">{merges}</mergeCells>',
    )
    _rejected(_workbook(sheet), "merged_ranges")


def test_merged_ranges_within_limits_still_fill_from_anchor():
    sheet = _sheet(
        '<row r="1"><c r="A1" t="inlineStr"><is><t>공급자</t></is></c></row>'
        '<row r="2"><c r="A2"><v>1</v></c><c r="B2"><v>2</v></c></row>',
        merges='<mergeCells count="1"><mergeCell ref="A1:B1"/></mergeCells>',
    )
    (table,) = parse_xlsx(_workbook(sheet)).tables
    assert table.rows[0].texts() == ["공급자", "공급자"]
    assert table.rows[0].cells[1].locator["merged"] == "A1:B1"


# ------------------------------------------------------------------ container / XML parser
def test_zip_bomb_by_overall_ratio(no_openpyxl):
    # each member stays under the 1 MiB per-entry ratio threshold; together ~50 MB of zeros
    extra = {f"xl/media/pad{i}.bin": b"\0" * 1_000_000 for i in range(50)}
    content = _workbook(_sheet('<row r="1"><c r="A1"><v>1</v></c></row>'), **extra)
    assert len(content) < 200_000
    t0 = time.perf_counter()
    with pytest.raises(IngestRejected) as e:
        parse_bytes(content, "bomb.xlsx")
    assert e.value.code == "zip_bomb"
    assert time.perf_counter() - t0 < BOUND_S


def test_zip_bomb_single_member_ratio(no_openpyxl):
    sheet = _sheet('<row r="1"><c r="A1"><v>1</v></c></row>' + " " * 5_000_000)
    _rejected(_workbook(sheet), "zip_bomb")


def test_entity_expansion_in_shared_strings_is_refused(no_openpyxl):
    lol = (
        '<?xml version="1.0"?><!DOCTYPE sst [<!ENTITY a "aaaaaaaaaa">'
        + "".join(f'<!ENTITY {chr(98 + i)} "{("&" + chr(97 + i) + ";") * 10}">' for i in range(8))
        + f']><sst xmlns="{NS}" count="1" uniqueCount="1"><si><t>&i;</t></si></sst>'
    ).encode()
    _rejected(_workbook(_sheet("<row/>"), **{"xl/sharedStrings.xml": lol}), "xml_dtd")


def test_too_many_sheet_entries_in_workbook(no_openpyxl):
    src = zipfile.ZipFile(io.BytesIO(_base()))
    wbxml = src.read("xl/workbook.xml").decode()
    extra = "".join(
        f'<sheet xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
        f' name="S{i}" sheetId="{i + 10}" r:id="rId1"/>'
        for i in range(60)
    )
    wbxml = wbxml.replace("</sheets>", extra + "</sheets>")
    _rejected(_replace(_base(), {"xl/workbook.xml": wbxml.encode()}), "sheets")


# ------------------------------------------------------------------ statuses
def test_rejections_are_failed_or_corrupt_never_empty_parsed(no_openpyxl):
    far = _workbook(_sheet('<row r="1048576"><c r="XFD1048576"><v>1</v></c></row>'))
    doc = safe_parse(far, "far.xlsx")
    assert doc.status is DocumentStatus.FAILED and "rows" in (doc.reason or "")
    assert doc.meta.get("error_code") == "rows" and not doc.tables

    bad_ref = _workbook(_sheet('<row r="1"><c r="AAAA1"><v>1</v></c></row>'))
    doc = safe_parse(bad_ref, "bad.xlsx")
    assert doc.status is DocumentStatus.CORRUPT and not doc.tables
    with pytest.raises(CorruptFile):
        parse_xlsx(bad_ref)

    broken = _workbook(b'<worksheet xmlns="' + NS.encode() + b'"><sheetData><row>')
    assert safe_parse(broken, "broken.xlsx").status is DocumentStatus.CORRUPT


def test_damage_found_while_streaming_cells_is_corrupt():
    # a shared-string reference with no shared-string table: read-only openpyxl only
    # notices while iterating, after load_workbook succeeded
    content = _workbook(_sheet('<row r="1"><c r="A1" t="s"><v>7</v></c></row>'))
    doc = safe_parse(content, "lazy.xlsx")
    assert doc.status is DocumentStatus.CORRUPT and not doc.tables


# ------------------------------------------------------------------ parts found like openpyxl
# openpyxl finds the workbook and shared-string parts through content types and
# relationships, never through the root element's name, and libxml2/expat read UTF-16
# without a byte order mark. The guard must bound the same parts in the same decoding.
SST_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"


def _with_shared_strings(sst_xml: bytes) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(_base()))
    ct = src.read("[Content_Types].xml").decode()
    ct = ct.replace(
        "</Types>", f'<Override PartName="/xl/sharedStrings.xml" ContentType="{SST_CT}"/></Types>'
    )
    sheet = _sheet('<row r="1"><c r="A1" t="s"><v>0</v></c></row>')
    return _replace(
        _base(),
        {"[Content_Types].xml": ct.encode(), SHEET: sheet, "xl/sharedStrings.xml": sst_xml},
    )


def _sst(root: str, n: int, *, decl: str = "UTF-8") -> str:
    items = "".join(f"<si><t>s{i}</t></si>" for i in range(n))
    return f'<?xml version="1.0" encoding="{decl}"?><{root} xmlns="{NS}">{items}</{root}>'


def test_shared_string_cap_holds_when_the_root_element_is_renamed(no_openpyxl):
    small = XlsxLimits(max_shared_strings=100)
    _rejected(_with_shared_strings(_sst("sst", 500).encode()), "shared_strings", xlsx_limits=small)
    renamed = _with_shared_strings(_sst("sstX", 500).encode())
    _rejected(renamed, "shared_strings", xlsx_limits=small)


def test_one_shared_string_with_too_many_runs(no_openpyxl):
    runs = "".join(f"<r><t>{i}</t></r>" for i in range(600))
    sst = f'<sst xmlns="{NS}"><si>{runs}</si></sst>'.encode()
    small = XlsxLimits(max_string_elements=500)
    _rejected(_with_shared_strings(sst), "shared_strings", xlsx_limits=small)


def test_utf16_without_bom_is_refused_but_with_bom_is_read():
    no_bom = _sst("sst", 3, decl="UTF-16").encode("utf-16-le")
    _rejected(_with_shared_strings(no_bom), "xml_encoding")
    no_bom_be = _sst("sst", 3, decl="UTF-16").encode("utf-16-be")
    _rejected(_with_shared_strings(no_bom_be), "xml_encoding")
    # UTF-16 with a byte order mark is valid OOXML and both parsers agree on it
    with_bom = _sst("sst", 3, decl="UTF-16").encode("utf-16")
    doc = parse_xlsx(_with_shared_strings(with_bom))
    assert doc.tables[0].rows[0].cells[0].text == "s0"


def test_dtd_is_refused_in_every_decoding(no_openpyxl):
    body = (
        '<!DOCTYPE sst [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;&a;">]>'
        f'<sst xmlns="{NS}"><si><t>&b;</t></si></sst>'
    )
    decl = '<?xml version="1.0" encoding="UTF-16"?>'
    _rejected(_with_shared_strings((decl + body).encode("utf-16")), "xml_dtd")
    _rejected(_with_shared_strings((decl + body).encode("utf-16-le")), "xml_encoding")
    _rejected(_with_shared_strings(("\n  " + decl + body).encode("utf-16-le")), "xml_encoding")


def test_declared_encodings_other_than_utf8_are_refused(no_openpyxl):
    sst = f'<sst xmlns="{NS}"><si><t>a</t></si></sst>'
    latin = '<?xml version="1.0" encoding="ISO-8859-1"?>' + sst
    _rejected(_with_shared_strings(latin.encode("latin-1")), "xml_encoding")
    utf7 = '<?xml version="1.0" encoding="UTF-7"?>' + sst
    _rejected(_with_shared_strings(utf7.encode()), "xml_encoding")


def test_sheet_entry_cap_holds_when_the_workbook_root_is_renamed(no_openpyxl):
    src = zipfile.ZipFile(io.BytesIO(_base()))
    wbxml = src.read("xl/workbook.xml").decode()
    extra = "".join(
        f'<sheet xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
        f' name="S{i}" sheetId="{i + 10}" r:id="rId1"/>'
        for i in range(60)
    )
    wbxml = wbxml.replace("</sheets>", extra + "</sheets>")
    wbxml = wbxml.replace("<workbook ", "<workbookX ").replace("</workbook>", "</workbookX>")
    assert "<workbookX " in wbxml
    _rejected(_replace(_base(), {"xl/workbook.xml": wbxml.encode()}), "sheets")


def test_styles_bomb_is_rejected_before_openpyxl(no_openpyxl):
    xf = '<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
    styles = (
        f'<styleSheet xmlns="{NS}"><fonts count="1"><font/></fonts>'
        f'<cellXfs count="3000">{xf * 3000}</cellXfs></styleSheet>'
    ).encode()
    content = _replace(_base(), {"xl/styles.xml": styles})
    _rejected(content, "styles", xlsx_limits=XlsxLimits(max_style_records=1000))
    # default caps: a 0.1 MiB upload with 200 000 <xf> took seconds and hundreds of MiB
    # varied ids keep the compression ratio realistic, so only the style cap can stop it
    xfs = "".join(f'<xf numFmtId="{i % 200}" fontId="{i % 7}"/>' for i in range(200_000))
    bomb = f'<styleSheet xmlns="{NS}"><cellXfs>{xfs}</cellXfs></styleSheet>'.encode()
    content = _replace(_base(), {"xl/styles.xml": bomb})
    assert len(content) < 1024 * 1024
    _rejected(content, "styles")


def test_manifest_and_relationship_parts_have_element_caps(no_openpyxl):
    src = zipfile.ZipFile(io.BytesIO(_base()))
    ct = src.read("[Content_Types].xml").decode()
    filler = "".join(
        f'<Override PartName="/x/{i}.xml" ContentType="application/xml"/>' for i in range(400)
    )
    small = XlsxLimits(max_part_elements=300)
    big_ct = ct.replace("</Types>", filler + "</Types>").encode()
    _rejected(_replace(_base(), {"[Content_Types].xml": big_ct}), "xml_elements", xlsx_limits=small)
    rel = (
        '<Relationship Id="h{0}" TargetMode="External" Target="https://example.invalid/{0}"'
        ' Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"/>'
    )
    rels = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + "".join(rel.format(i) for i in range(400))
        + "</Relationships>"
    ).encode()
    # openpyxl parses each sheet's relationships even in read-only mode
    sheet_rels = _replace(_base(), {"xl/worksheets/_rels/sheet1.xml.rels": rels})
    _rejected(sheet_rels, "xml_elements", xlsx_limits=small)


def test_part_openpyxl_parses_but_is_not_xml_fails_closed(no_openpyxl):
    with pytest.raises(CorruptFile):
        parse_xlsx(_replace(_base(), {"xl/styles.xml": b"\x89PNG\r\n\x1a\n" + b"\x00" * 64}))


def test_binary_members_with_leading_nul_bytes_are_not_mistaken_for_xml():
    # EMF images start with 01 00 00 00; they are opened, never parsed as XML
    emf = b"\x01\x00\x00\x00" + bytes(range(256)) * 4
    doc = parse_xlsx(_replace(_base(), {"xl/media/image1.emf": emf}))
    assert doc.tables[0].rows[0].cells[0].text == "x"


def test_workbook_with_chart_sheet_and_number_formats_still_parses():
    from openpyxl.chart import BarChart, Reference

    from jettae.ingest.xlsx_guard import inspect_xlsx

    wb = Workbook()
    ws = wb.active
    for r in range(1, 6):
        ws.cell(r, 1, r * 1000).number_format = "#,##0"
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=1, min_row=1, max_row=5))
    wb.create_chartsheet("chart").add_chart(chart)
    buf = io.BytesIO()
    wb.save(buf)
    content = buf.getvalue()
    assert any(n.startswith("xl/charts/") for n in zipfile.ZipFile(buf).namelist())
    scan = inspect_xlsx(content, XlsxLimits())
    assert scan.aux_elements > 0 and len(scan.sheets) == 1
    doc = parse_xlsx(content)
    assert [t.name for t in doc.tables] == [ws.title]
    assert doc.tables[0].rows[4].cells[0].raw == 5000
