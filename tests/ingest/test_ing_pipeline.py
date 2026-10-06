from datetime import UTC, date, datetime

import pytest
from ing_helpers import TENANT, save_wb, scanned_pdf
from openpyxl import Workbook

from jettae.app import JettaeService, in_memory_repositories
from jettae.domain.models import Change
from jettae.domain.money import Money
from jettae.domain.status import ChangeKind, DocKind, DocumentStatus, ReconcileStatus
from jettae.ingest import IngestOptions, ingest_document
from jettae.verify.checks import check_citation

NOW = datetime(2025, 11, 1, tzinfo=UTC)
OPTS = IngestOptions(counterparty_override="가나유통")


@pytest.fixture
def svc():
    return JettaeService(in_memory_repositories(), clock=lambda: NOW)


def settlement(rows):
    wb = Workbook()
    ws = wb.active
    ws.append(["정산번호", "거래형태", "상품수령일", "판매마감일", "정산금액"])
    for r in rows:
        ws.append(r)
    return save_wb(wb)


BANK = (
    "거래일자,적요,보낸분/받는분,출금액,입금액,잔액\n"
    '2025-10-20,타행입금,가나유통,0,"10,000,000","10,000,000"\n'
).encode("cp949")

S1 = [
    ["A-1", "직매입", "2025-08-07", None, 10_000_000],
    ["A-2", "직매입", None, None, 2_000_000],
]


def test_ingest_registers_document_records_and_citable_facts(svc):
    doc, res, changes = ingest_document(
        svc, TENANT, filename="정산서.xlsx", content=settlement(S1), document_id="s1", options=OPTS
    )
    assert doc.status is DocumentStatus.PARSED and doc.kind is DocKind.SETTLEMENT
    assert doc.media_type.endswith("spreadsheetml.sheet")
    lines = svc.repos.ledger.list(TENANT, "settlement_line")
    assert len(lines) == 2
    assert all(c.kind is ChangeKind.ADD for c in changes)
    text = svc.repos.documents.get_text(TENANT, doc.id)
    stored = svc.repos.facts.list(TENANT)
    assert stored and all(check_citation(f.span, text).passed for f in stored if f.span is not None)
    assert all(f.span.doc_version_id == doc.id for f in stored if f.span is not None)


def test_analysis_on_ingested_records_and_missing_base_date(svc):
    ingest_document(
        svc, TENANT, filename="s.xlsx", content=settlement(S1), document_id="s1", options=OPTS
    )
    ingest_document(svc, TENANT, filename="bank.csv", content=BANK, document_id="b1")
    svc.run_analysis(TENANT, as_of=date(2025, 11, 1))
    views = {v.decision.subject_id: v.decision for v in svc.list_decisions(TENANT)}
    lines = {ln.settlement_ref: ln for ln in svc.repos.ledger.list(TENANT, "settlement_line")}
    matched = views[lines["A-1"].id]
    assert matched.status is ReconcileStatus.MATCHED
    pending = views[lines["A-2"].id]
    assert "goods_received_date" in pending.missing  # never filled from another date
    assert pending.required_documents


def test_new_version_updates_and_removes_rows(svc):
    d1, _, _ = ingest_document(
        svc, TENANT, filename="s.xlsx", content=settlement(S1), document_id="s1", options=OPTS
    )
    svc.run_analysis(TENANT, as_of=date(2025, 11, 1))
    corrected = [["A-1", "직매입", "2025-08-07", None, 10_000_000]]  # A-2 withdrawn
    v2_bytes = settlement(corrected)  # built once: xlsx embeds a save timestamp
    d2, res2, changes = ingest_document(
        svc,
        TENANT,
        filename="s_v2.xlsx",
        content=v2_bytes,
        document_id="s1",
        options=OPTS,
    )
    assert d2.version == 2 and d2.supersedes == d1.id
    kinds = {(c.entity, c.kind) for c in changes}
    assert ("settlement_line", ChangeKind.UPDATE) in kinds  # A-1 kept, new spans
    assert ("settlement_line", ChangeKind.REMOVE) in kinds  # A-2 removed
    lines = svc.repos.ledger.list(TENANT, "settlement_line")
    assert [ln.settlement_ref for ln in lines] == ["A-1"]
    facts = svc.repos.facts.list(TENANT)
    assert all(f.span.doc_version_id == d2.id for f in facts if f.span is not None)
    assert not any(f.subject_id not in {lines[0].id} for f in facts)
    # identical re-upload: same version, only UPDATE changes (nothing removed)
    d3, _, ch3 = ingest_document(
        svc,
        TENANT,
        filename="s_v2.xlsx",
        content=v2_bytes,
        document_id="s1",
        options=OPTS,
    )
    assert d3.id == d2.id and all(c.kind is ChangeKind.UPDATE for c in ch3)


def test_scanned_pdf_registered_as_unsupported_scan_without_records(svc):
    doc, res, changes = ingest_document(svc, TENANT, filename="scan.pdf", content=scanned_pdf())
    assert doc.status is DocumentStatus.UNSUPPORTED_SCAN
    assert changes == [] and res.records == []


def test_needs_mapping_registers_without_records_then_user_mapping(svc):
    wb = Workbook()
    ws = wb.active
    ws.append(["번호", "형태", "받은날", "값"])
    ws.append(["A-1", "직매입", "2025-08-07", "10,000,000"])
    raw = save_wb(wb)
    doc, res, changes = ingest_document(
        svc, TENANT, filename="x.xlsx", content=raw, document_id="x"
    )
    assert doc.status is DocumentStatus.NEEDS_MAPPING and changes == []
    opts = IngestOptions(
        counterparty_override="가나유통",
        format_id="retail_settlement",
        mapping={"*": {"amount": "값", "trade_type": "형태", "goods_received_date": "받은날"}},
    )
    doc2, res2, ch2 = ingest_document(
        svc, TENANT, filename="x.xlsx", content=raw, document_id="x", options=opts
    )
    assert doc2.status is DocumentStatus.NEEDS_MAPPING  # same content → same registered version
    (line,) = res2.records
    assert line.amount == Money(10_000_000) and line.goods_received_date == date(2025, 8, 7)
    m = res2.plan.tables[0].mapping.match("goods_received_date")
    assert m is not None and m.confidence.value == "confirmed"
    assert isinstance(ch2[0], Change)
