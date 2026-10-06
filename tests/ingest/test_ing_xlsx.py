from datetime import date, datetime

import pytest
from ing_helpers import TENANT, add_zip_member, save_wb
from openpyxl import Workbook

from jettae.domain.models import Invoice, SettlementLine
from jettae.domain.money import Money
from jettae.domain.status import DocKind, DocumentStatus, LineKind, TradeType
from jettae.ingest import IngestOptions, IngestRejected, ingest_bytes, parse_bytes, safe_parse
from jettae.ingest.xlsx import parse_xlsx
from jettae.verify.checks import check_citation


def hometax_two_row_header() -> bytes:
    """Title + two-row header with merged group cells (공급자 / 공급받는자), serial dates."""
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "안내"
    ws0["A1"] = "이 파일은 테스트용 손 예제입니다"
    ws = wb.create_sheet("세금계산서")
    ws["A1"] = "매출 전자세금계산서 목록"
    ws.merge_cells("A1:J1")
    hdr1 = [
        "작성일자",
        "승인번호",
        "공급자",
        None,
        None,
        "공급받는자",
        None,
        None,
        "합계금액",
        "공급가액",
    ]
    hdr2 = [None, None, "등록번호", "상호", "성명", "등록번호", "상호", "성명", None, None]
    for c, v in enumerate(hdr1, start=1):
        ws.cell(row=3, column=c, value=v)
    for c, v in enumerate(hdr2, start=1):
        ws.cell(row=4, column=c, value=v)
    ws.merge_cells("C3:E3")
    ws.merge_cells("F3:H3")
    for col in "ABIJ":
        ws.merge_cells(f"{col}3:{col}4")
    ws.cell(row=3, column=11, value="세액")
    ws.merge_cells("K3:K4")
    rows = [
        # 작성일자 as a General-format Excel serial number (45876 = 2025-08-07)
        (
            45876,
            "20250807-41000000-12345678",
            "123-45-67890",
            "제때상사",
            "김가나",
            "220-81-00001",
            "(주)가나유통",
            "이대표",
            11_000_000,
            10_000_000,
            1_000_000,
        ),
        (
            datetime(2025, 8, 31),
            "20250831-41000000-00000001",
            "123-45-67890",
            "제때상사",
            "김가나",
            "220-81-00001",
            "(주)가나유통",
            "이대표",
            -1_100_000,
            -1_000_000,
            -100_000,
        ),
    ]
    for r, vals in enumerate(rows, start=5):
        for c, v in enumerate(vals, start=1):
            ws.cell(row=r, column=c, value=v)
    ws.cell(row=7, column=1, value="합계")
    ws.cell(row=7, column=9, value=9_900_000)
    return save_wb(wb)


def test_merged_two_row_header_is_composed_and_cells_keep_locators():
    doc = parse_xlsx(hometax_two_row_header(), filename="ht.xlsx")
    sheet = next(t for t in doc.tables if t.name == "세금계산서")
    r3 = next(r for r in sheet.rows if r.index == 3)
    # merged group label copied to every covered cell, with the merged range in the locator
    assert r3.cells[3].text == "공급자" and r3.cells[3].locator["merged"] == "C3:E3"
    assert r3.cells[3].locator["cell"] == "D3"
    plan = safe_parse(hometax_two_row_header(), "ht.xlsx")
    res = ingest_bytes(hometax_two_row_header(), "ht.xlsx", tenant_id=TENANT, doc_version_id="dv")
    assert plan.ok
    tp = next(t for t in res.plan.tables if t.table == "세금계산서")
    assert tp.recognition.view.header_rows == (3, 4)
    assert "공급받는자 상호" in tp.recognition.view.headers
    assert tp.spec is not None and tp.spec.id == "hometax_etax_list"
    m = {x.field: x for x in tp.mapping.matches}
    assert m["buyer_name"].header == "공급받는자 상호"
    assert m["supplier_name"].header == "공급자 상호"
    # the 안내 sheet is ignored, not an error
    other = next(t for t in res.plan.tables if t.table == "안내")
    assert not other.relevant


def test_hometax_rows_become_invoices_without_base_dates():
    res = ingest_bytes(hometax_two_row_header(), "ht.xlsx", tenant_id=TENANT, doc_version_id="dv")
    assert res.status is DocumentStatus.PARSED
    assert res.plan.doc_kind is DocKind.TAX_INVOICE
    inv = res.records
    assert len(inv) == 2 and all(isinstance(i, Invoice) for i in inv)
    a, b = inv
    assert a.issue_date == date(2025, 8, 7)  # from the Excel serial number
    assert a.reference == "20250807-41000000-12345678"
    assert a.counterparty == "(주)가나유통"  # "매출" title → we are the supplier
    assert a.amount == Money(11_000_000)
    assert b.amount == Money(-1_100_000)  # 수정(마이너스) 세금계산서 kept negative
    # tax-invoice date is never a base date
    assert a.goods_received_date is None and a.sales_close_date is None and a.trade_type is None
    assert {"trade_type", "goods_received_date", "sales_close_date"} <= a.missing
    serial_fact = next(f for f in res.facts if f.id == f"fact:{a.id}:issue_date")
    assert serial_fact.span.locator["note"] == "excel_serial"
    assert serial_fact.span.locator["sheet"] == "세금계산서"
    assert serial_fact.span.locator["cell"] == "A5"
    # 합계 row: stated 9,900,000 equals 11,000,000 - 1,100,000
    (tc,) = [t for t in res.totals if t.field == "total"]
    assert tc.matches and tc.row == 7
    for f in res.facts:
        if f.span is not None:
            assert check_citation(f.span, res.plan.parsed.text).passed


def flat_hometax(title: str = "") -> bytes:
    wb = Workbook()
    ws = wb.active
    if title:
        ws.append([title])
    ws.append(
        [
            "작성일자",
            "승인번호",
            "발급일자",
            "전송일자",
            "공급자사업자등록번호",
            "종사업장번호",
            "상호",
            "대표자명",
            "공급받는자사업자등록번호",
            "종사업장번호",
            "상호",
            "대표자명",
            "합계금액",
            "공급가액",
            "세액",
            "전자세금계산서분류",
            "전자세금계산서종류",
            "발급유형",
            "비고",
            "영수/청구 구분",
        ]
    )
    ws.append(
        [
            "2025-09-01",
            "20250901-10000000-00000077",
            "2025-09-01",
            "2025-09-02",
            "1234567890",
            "",
            "제때상사",
            "김가나",
            "2208100001",
            "",
            "다라마트",
            "박",
            "5,500,000",
            "5,000,000",
            "500,000",
            "세금계산서",
            "일반",
            "인터넷발급",
            "",
            "청구",
        ]
    )
    return save_wb(wb)


def test_flat_layout_repeated_party_columns_and_direction_from_self_brn():
    raw = flat_hometax()
    # no 매출/매입 hint: direction unknown → counterparty not guessed
    res = ingest_bytes(raw, "list.xlsx")
    (inv,) = res.records
    assert inv.counterparty == "" and {"counterparty", "direction"} <= inv.missing
    # with our BRN the row is a sales invoice
    res2 = ingest_bytes(raw, "list.xlsx", options=IngestOptions(self_brn="123-45-67890"))
    (inv2,) = res2.records
    assert inv2.counterparty == "다라마트" and inv2.amount == Money(5_500_000)
    brn = next(f for f in res2.facts if f.kind == "invoice.buyer_brn")
    assert brn.value == "220-81-00001"
    # a purchase list produces no receivables
    res3 = ingest_bytes(raw, "list.xlsx", options=IngestOptions(self_brn="220-81-00001"))
    assert res3.records == [] and any("매입" in i.message for i in res3.issues)


def settlement_wb(with_formula: bool = False) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "정산"
    ws.append(["가나유통 9월 정산서"])
    ws.append([])
    ws.append(
        ["정산번호", "거래형태", "구분", "상품수령일", "판매마감일", "정산금액", "지급예정일"]
    )
    ws.append(["00123", "직매입", "매입", datetime(2025, 8, 7), None, 10_000_000, "2025-10-20"])
    ws.append([124, "특약매입", "판매", None, 45930, 3_000_000, "2025-10-20"])
    ws.append([125, "특약매입", "판촉비 공제", None, "2025-09-30", -200_000, None])
    ws.append([126, "", "", None, None, 500_000, None])
    ws["A5"].number_format = "00000"  # stored as number, displayed with leading zeros
    ws["A6"].number_format = "00000"
    ws["A7"].number_format = "00000"
    if with_formula:
        ws.append(["합계", None, None, None, None, "=SUM(F4:F7)", None])
    return save_wb(wb)


def test_settlement_sheet_base_dates_only_from_their_own_columns():
    res = ingest_bytes(
        settlement_wb(),
        "s.xlsx",
        options=IngestOptions(counterparty_override="가나유통"),
        doc_version_id="dv",
    )
    assert res.status is DocumentStatus.PARSED
    lines = res.records
    assert all(isinstance(x, SettlementLine) for x in lines)
    a, b, c, d = lines
    assert a.settlement_ref == "00123" and b.settlement_ref == "00124"  # leading zeros kept
    assert a.trade_type is TradeType.DIRECT and a.goods_received_date == date(2025, 8, 7)
    assert a.missing == frozenset()
    assert b.trade_type is TradeType.CONSIGNMENT and b.sales_close_date == date(2025, 9, 30)
    assert c.line_kind is LineKind.DEDUCTION and c.amount == Money(-200_000)
    # no trade type / no base date: nothing inferred from 지급예정일
    assert d.trade_type is None and d.goods_received_date is None and d.sales_close_date is None
    assert {"trade_type", "goods_received_date", "sales_close_date"} <= d.missing
    assert all(x.counterparty == "가나유통" for x in lines)
    pay = [f for f in res.facts if f.kind == "settlement_line.payment_date"]
    assert pay and all(f.value == date(2025, 10, 20) for f in pay)


def test_formula_cells_are_not_evaluated():
    raw = settlement_wb(with_formula=True)
    doc = parse_xlsx(raw)
    t = doc.tables[0]
    total_row = next(r for r in t.rows if r.index == 8)
    f_cell = total_row.cells[5]
    # openpyxl-written files have no cached value: the cell stays empty, flagged as formula
    assert f_cell.locator.get("formula") is True and f_cell.text == ""
    assert any("formula" in w for w in doc.warnings)
    res = ingest_bytes(raw, "s.xlsx", options=IngestOptions(counterparty_override="가나유통"))
    assert len(res.records) == 4  # 합계 row is a total row, not a line


def test_macro_and_binary_parts_rejected():
    raw = add_zip_member(settlement_wb(), "xl/vbaProject.bin", b"\x00" * 10)
    with pytest.raises(IngestRejected) as ei:
        parse_bytes(raw, "m.xlsx")
    assert ei.value.code == "macro"
    doc = safe_parse(raw, "m.xlsm")
    assert doc.status is DocumentStatus.FAILED and "macro" in (doc.reason or "")


def test_corrupt_xlsx_is_corrupt_not_empty():
    raw = settlement_wb()[:300]
    doc = safe_parse(raw, "broken.xlsx")
    assert doc.status is DocumentStatus.CORRUPT
    assert doc.reason
